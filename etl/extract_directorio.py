"""Extrae SOLO secretarios y directores (nombre, cargo, celular) de la hoja 'Gabinete' del directorio protocolario.
No toma cédulas, correos, subsecretarios, asesores ni jefes de oficina."""
import openpyxl, re, csv
from clean import canon_one, fold

def title(s):
    PART={'de','del','la','las','los','y','e','en'}; ACR={'dagma':'DAGMA','dadii':'DADII','datic':'DATIC','uaespm':'UAESPM','uae':'UAE','tic':'TIC'}
    out=[]
    for i,w in enumerate(re.sub(r'\s+',' ',s).strip().lower().split(' ')):
        out.append(ACR.get(w) or (w if (w in PART and i>0) else w[:1].upper()+w[1:]))
    return ' '.join(out)

def phone(t):
    """Primer celular colombiano del texto, en formato 57XXXXXXXXXX."""
    if t is None: return ''
    t=re.sub(r'\.0\b','',str(t).replace('\n',' '))
    m=re.search(r'(?<!\d)(3\d{2})[\s\-]*(\d{7})(?!\d)',t)
    return '57'+m.group(1)+m.group(2) if m else ''

def extraer(path):
    ws=openpyxl.load_workbook(path,data_only=True)['Gabinete']
    out=[]
    for r in range(1,ws.max_row+1):
        a=ws.cell(r,1).value
        if isinstance(a,str) and 'ENTIDADES DESCENTRALIZADAS' in a.upper(): break   # fin del gabinete central
        nombre=ws.cell(r,3).value; cargo=ws.cell(r,4).value
        if not isinstance(cargo,str): continue
        f=fold(cargo)
        if not re.search(r'secretari|direct',f) or re.search(r'sub|asesor|jefe|gerente',f): continue
        org=canon_one(cargo)
        if not org: continue
        out.append(dict(organismo=org[0],cargo=title(cargo),
                        nombre=title(re.sub(r'\s+',' ',nombre.replace('\n',' '))) if isinstance(nombre,str) and nombre.strip() else '',
                        telefono=phone(ws.cell(r,6).value),fila_origen=r))
    return out

def guardar_csv(rows,path):
    with open(path,'w',newline='',encoding='utf-8') as f:
        w=csv.DictWriter(f,fieldnames=['organismo','cargo','nombre','telefono']); w.writeheader()
        for x in rows: w.writerow({k:x[k] for k in w.fieldnames})
