import openpyxl, pandas as pd, re, datetime as dt
# Rutas de los Excel de control. Se sobrescriben desde run_etl.py (ver README).
FILES={2024:'data/control_2024.xlsx',2025:'data/control_2025.xlsx',2026:'data/control_2026.xlsx'}
def clean(v):
    if v is None: return None
    if isinstance(v,str):
        v=re.sub(r'\s+',' ',v.replace('\xa0',' ')).strip()
        return v if v not in ('','-') else None
    return v
def load(vig):
    wb=openpyxl.load_workbook(FILES[vig],data_only=True); ws=wb['CONTROL']; rows=[]
    for i,r in enumerate(ws.iter_rows(min_row=6,values_only=True),6):
        r=[clean(c) for c in r]
        if vig<2026:
            if not any(r[1:15]): continue
            rows.append(dict(vigencia=vig,fila_excel=i,orden=r[0],rad_ente=r[1],fecha_oficio=r[2],ente=r[3],orfeo_entrada=r[4],fecha_radicacion=r[5],fecha_daci=r[6],dirigido=r[7],asunto=r[8],dias_habiles=r[9],fecha_max=r[10],responde=r[11],orfeo_resp=r[12],fecha_resp=r[13],estado=r[14],obs=r[15],clase=None))
        else:
            if not any(r[1:13]): continue
            rows.append(dict(vigencia=vig,fila_excel=i,orden=r[0],rad_ente=r[1],fecha_oficio=None,ente=r[2],orfeo_entrada=r[3],fecha_radicacion=r[4],fecha_daci=None,dirigido=r[6],asunto=r[7],dias_habiles=r[8],fecha_max=r[9],responde=None,orfeo_resp=r[10],fecha_resp=r[11],estado=r[12],obs=r[13],clase=r[5]))
    return pd.DataFrame(rows)
def load_norespuesta(vig):
    wb=openpyxl.load_workbook(FILES[vig],data_only=True); ws=wb['NO RESPUESTA']; rows=[]
    for i,r in enumerate(ws.iter_rows(min_row=6,values_only=True),6):
        r=[clean(c) for c in r]
        if not any(r[1:9]): continue
        rows.append(dict(vigencia=vig,fila_excel=i,orden=r[0],rad_ente=r[1],fecha_oficio=r[2],ente=r[3],orfeo_entrada=r[4],fecha_radicacion=r[5],fecha_daci=r[6],dirigido=r[7],asunto=r[8],traslado=r[9] if len(r)>9 else None))
    return pd.DataFrame(rows)
